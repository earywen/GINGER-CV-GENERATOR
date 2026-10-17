# -*- coding: utf-8 -*-
import io
import os
import subprocess
import tempfile
import shutil
from typing import Sequence, Optional
import pypdf
import pdfplumber
import docx
import pptx

from app.core import get_logger

logger = get_logger(__name__)

def extract_text_with_anymd(file_bytes: bytes, ext: str = ".pdf") -> Optional[str]:
    """
    Extracts structured Markdown from file bytes using anymd CLI.
    Preserves multi-column layout, markdown tables, and supports OCR on scanned documents.
    """
    cmd = None
    if shutil.which("anymd"):
        cmd = ["anymd"]
    elif shutil.which("npx"):
        cmd = ["npx", "--yes", "@sylphx/anymd"]

    if not cmd:
        return None

    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name

    try:
        proc = subprocess.run(
            cmd + [tmp_path],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=15
        )
        if proc.returncode == 0:
            out = proc.stdout.strip()
            if len(out) > 50:
                return out
    except Exception as e:
        logger.warning(f"anymd extraction error: {e}")
    finally:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass
    return None

def _format_table_row(cells: Sequence) -> str:
    """Formats table row cells as pipe-separated text."""
    row_text = " | ".join(cell.text.strip() for cell in cells if cell.text.strip())
    return row_text

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extracts text from PDF using pdfplumber with a clean fallback to pypdf.
    """
    # Attempt 1: pdfplumber (best for layout and text positioning)
    plumber_parts = []
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    plumber_parts.append(extracted.strip())
    except Exception as e:
        logger.warning(f"pdfplumber extraction error: {e}")

    joined_plumber = "\n".join(plumber_parts).strip()
    if plumber_parts and len(joined_plumber) > 50:
        return "\n\n".join(plumber_parts)

    # Attempt 2: pypdf (fallback on fresh buffer if pdfplumber extracted nothing or failed)
    pypdf_parts = []
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            t = page.extract_text()
            if t:
                pypdf_parts.append(t.strip())
    except Exception as e:
        logger.warning(f"pypdf extraction error: {e}")

    joined_pypdf = "\n".join(pypdf_parts).strip()
    if pypdf_parts and len(joined_pypdf) > 50:
        return "\n\n".join(pypdf_parts)

    # If both extracted minimal content, return whichever has more text
    if len(joined_plumber) >= len(joined_pypdf):
        return "\n\n".join(plumber_parts) if plumber_parts else ""
    return "\n\n".join(pypdf_parts) if pypdf_parts else ""

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text from OpenXML Word (.docx) document including tables."""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = _format_table_row(row.cells)
            if row_text:
                paragraphs.append(row_text)
    return "\n\n".join(paragraphs)

def extract_text_from_pptx(file_bytes: bytes) -> str:
    """Extracts text from OpenXML PowerPoint (.pptx) document including tables."""
    prs = pptx.Presentation(io.BytesIO(file_bytes))
    text_runs = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in paragraph.runs).strip()
                    if t:
                        text_runs.append(t)
            elif shape.has_table:
                for row in shape.table.rows:
                    row_text = _format_table_row(row.cells)
                    if row_text:
                        text_runs.append(row_text)
    return "\n\n".join(text_runs)

def extract_text(filename: str, file_bytes: bytes) -> str:
    """
    Main extraction dispatcher based on file extension.
    Raises ValueError with a clear user message for obsolete binary formats (.doc, .ppt).
    """
    ext = os.path.splitext(filename.lower())[1]
    
    if ext == ".doc":
        raise ValueError(
            "Le format binaire ancien Word (.doc) n'est pas supporté. "
            "Veuillez enregistrer votre document en format .docx ou l'exporter en .pdf."
        )
    if ext == ".ppt":
        raise ValueError(
            "Le format binaire ancien PowerPoint (.ppt) n'est pas supporté. "
            "Veuillez enregistrer votre présentation en format .pptx ou l'exporter en .pdf."
        )

    if ext == ".pdf":
        anymd_text = extract_text_with_anymd(file_bytes, ".pdf")
        if anymd_text:
            return anymd_text
        return extract_text_from_pdf(file_bytes)
    elif ext == ".docx":
        return extract_text_from_docx(file_bytes)
    elif ext == ".pptx":
        return extract_text_from_pptx(file_bytes)
    else:
        # Default text/markdown/plain decoding
        return file_bytes.decode("utf-8", errors="ignore")
